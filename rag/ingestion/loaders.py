"""Reading files into clean text, with headings marked as markdown ("#", "##", ...).

Supported formats: md, markdown, txt, html, pdf, docx, pptx, xlsx, csv.

Headings matter because chunking cuts on them. HTML and Word documents declare them;
for PDFs, and Word files written without heading styles, they are recovered from
font size and weight, and as a last resort from wording ("Capitolul II", "Art. 5").
"""
from __future__ import annotations

import codecs
import csv
import hashlib
import io
import json
import re
import unicodedata
from bisect import bisect_right
from dataclasses import asdict, dataclass, field
from pathlib import Path

import charset_normalizer
import docx
import openpyxl
import pptx
import pymupdf
from bs4 import BeautifulSoup
from markdownify import markdownify


@dataclass
class Document:
    doc_id: str      # stable id, derived from the source
    source: str      # "<folder>/<file>"
    text: str        # cleaned text, headings as markdown "#" lines
    metadata: dict = field(default_factory=dict)


def sha(s: str, n: int = 16) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:n]


# ------------------------------------------------------------------ cleaning

# Cedilla letters (ş ţ), common in PDFs, replaced by the correct comma-below ones
# (ș ț); otherwise BM25 sees "ştiinţă" and "știință" as different words.
CEDILLA_FIX = str.maketrans({"ş": "ș", "Ş": "Ș", "ţ": "ț", "Ţ": "Ț"})

# Invisible characters that split or glue words for a keyword index.
INVISIBLE = str.maketrans({
    "\xa0": " ", "\u202f": " ", "\u2007": " ", "\u2009": " ",  # fixed-width spaces
    "\u00ad": "",  # soft hyphen
    "\u200b": "", "\u200c": "", "\u200d": "",  # zero-width space / non-joiner / joiner
    "\ufeff": "",  # BOM in the middle of a file
})

# Typographic ligatures ("ﬁ" -> "fi"). Expanded by hand: NFKC would also rewrite
# fractions, superscripts and full-width characters.
LIGATURES = str.maketrans({"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi",
                           "ﬄ": "ffl", "ﬅ": "st", "ﬆ": "st", "œ": "oe", "Œ": "Oe"})


def clean_text(text: str, from_pdf: bool = False) -> str:
    """Normalize line endings, Unicode and whitespace."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = unicodedata.normalize("NFC", text)
    # A soft hyphen at a line break joins the word across the break. This must happen
    # before INVISIBLE deletes soft hyphens, or the word stays split in two.
    text = re.sub(r"\u00ad[ \t]*\n[ \t]*", "", text)
    text = text.translate(CEDILLA_FIX).translate(INVISIBLE).translate(LIGATURES)
    if from_pdf:
        text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # words hyphenated across lines
        text = re.sub(r"[ \t]+", " ", text)
    # Other formats keep their spacing: indentation matters in code blocks.
    text = re.sub(r"[ \t]+\n", "\n", text)            # trailing whitespace
    text = re.sub(r"\n{3,}", "\n\n", text)            # at most one empty line
    return text.strip()


def decode(raw: bytes, rel: str = "") -> str:
    """Decode a text file: by its BOM, else as UTF-8, else by charset-normalizer's guess.

    Never decodes with errors="ignore": a Romanian file in cp1250 would lose every
    diacritic without any error.
    """
    for bom, encoding in ((codecs.BOM_UTF8, "utf-8-sig"),
                          (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16"),
                          (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32")):
        if raw.startswith(bom):
            return raw.decode(encoding)
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass
    guess = charset_normalizer.from_bytes(raw).best()
    if guess is None:
        print(f"[warn] {rel}: undecodable bytes, falling back to cp1250 with replacements")
        return raw.decode("cp1250", errors="replace")
    print(f"[warn] {rel}: not UTF-8, decoded as {guess.encoding}")
    return str(guess)


def read_text_file(path: Path, rel: str = "") -> str:
    return decode(path.read_bytes(), rel or path.name)


def load_text(path: Path, rel: str) -> Document:
    text = clean_text(read_text_file(path, rel))
    fmt = "markdown" if path.suffix.lower() in {".md", ".markdown"} else "text"
    if fmt == "text":  # no markup, so headings can only be guessed from wording
        text = apply_pattern_headings(text)
    return Document(sha(rel), rel, text, {"format": fmt})


# Containers that usually hold a page's own content, most specific first.
HTML_MAIN = ["main", "article", "[role=main]", "#content", ".entry-content", "#main",
             ".post", ".content"]
HTML_BOILERPLATE = ["head", "script", "style", "noscript", "nav", "footer", "header",
                    "aside", "form", "iframe"]


def html_main_content(soup: BeautifulSoup):
    """The element holding the page's own content: the first known content container
    with at least a fifth of the page's text, else the whole body.

    A saved web page carries the site's menus, news and footer around the content;
    as text they turned into chunks that matched almost any question about the site.
    """
    total = len(soup.get_text(" ", strip=True)) or 1
    for selector in HTML_MAIN:
        element = soup.select_one(selector)
        if element is not None and len(element.get_text(" ", strip=True)) >= total / 5:
            return element
    return soup.body or soup


def drop_link_lists(root) -> None:
    """Remove lists made almost only of links: menus that are not marked up as <nav>."""
    for lst in root.find_all(["ul", "ol"]):
        if getattr(lst, "decomposed", False):
            continue   # inside a list removed already
        links = lst.find_all("a")
        text = len(lst.get_text(" ", strip=True))
        linked = sum(len(a.get_text(" ", strip=True)) for a in links)
        if len(links) >= 5 and text and linked / text > 0.8:
            lst.decompose()


def load_html(path: Path, rel: str) -> Document:
    soup = BeautifulSoup(read_text_file(path, rel), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(HTML_BOILERPLATE):
        tag.decompose()
    content = html_main_content(soup)
    drop_link_lists(content)
    # Links keep their text but lose their URLs, which only add noise to the index.
    md = markdownify(str(content), heading_style="ATX", strip=["a", "img"])
    return Document(sha(rel), rel, clean_text(md), {"format": "html", "title": title})


# --------------------------------------------------- headings from formatting
#
# A PDF only knows that some text is drawn larger or bolder than the rest; a Word file
# can also have titles typed in bold instead of a heading style. In both cases the
# hierarchy is recovered from font size and weight.

HEADING_MAX_CHARS = 120
HEADING_SIZE_RATIO = 1.15  # how much larger than the body text a heading must be


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def repeat_key(text: str) -> str:
    """Normalized line with digits masked, so "Page 3 of 12" matches "Page 4 of 12"."""
    return re.sub(r"\d+", "#", norm(text))


def body_size(sized_texts: list[tuple[str, float]]) -> float:
    """The body font size: the size with the most characters set in it."""
    weight: dict[float, int] = {}
    for text, size in sized_texts:
        weight[round(size, 1)] = weight.get(round(size, 1), 0) + len(text.strip())
    return max(weight, key=weight.get) if weight else 0.0


def looks_like_heading(text: str, size: float, emphasized: bool, body: float) -> bool:
    text = text.strip()
    if len(text) < 2 or len(text) > HEADING_MAX_CHARS or text.isdigit():
        return False
    if text[-1] in ".,;:":  # headings rarely end in sentence punctuation
        return False
    return size >= body * HEADING_SIZE_RATIO or (emphasized and size >= body)


def level_map(sizes: set[float]) -> dict[float, int]:
    """Rank the heading sizes: the largest becomes "#", the next "##", up to 6."""
    return {size: min(i + 1, 6) for i, size in enumerate(sorted(sizes, reverse=True))}


# ------------------------------------------------- headings from wording
#
# For documents set entirely in one font. Legal and administrative texts name their
# divisions ("Capitolul II", "Articolul 5"). Numbered lines must recur, and a heading
# may not end like a sentence, so numbered lists are not mistaken for headings.

HEADING_KEYWORDS = [  # from the widest division to the narrowest
    r"partea|titlul|cartea|part|title|book",
    r"capitolul|cap\.|chapter",
    r"sec[țt]iunea|section|subcapitolul",
    r"articolul|art\.|article|§",
    r"anexa|annex|appendix",
]
# Annexes sit beside the chapters, not below the articles.
ANNEX_RANK = len(HEADING_KEYWORDS) - 1
KEYWORD_PATTERNS = [re.compile(rf"^({kw})\s*[:.]?\s*([IVXLCDM]+|\d+)\b", re.I)
                    for kw in HEADING_KEYWORDS]
NUMBERED = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(\S.*)$")
MIN_NUMBERED = 2  # a single "1. something" is a list item, not a section


def pattern_rank(text: str) -> int | None:
    """How deep a line's wording puts it in the hierarchy (lower is wider), or None
    for ordinary text. Ranks are mapped to "#" levels afterwards."""
    text = text.strip()
    if len(text) < 2 or len(text) > HEADING_MAX_CHARS or text[-1] in ".,;:":
        return None
    for rank, pattern in enumerate(KEYWORD_PATTERNS):
        if pattern.match(text):
            return rank
    numbered = NUMBERED.match(text)
    if numbered and numbered.group(2)[0].isupper():
        # "2.1 Definitions" sits one level under "2. Scope"
        return 10 + min(numbered.group(1).count(".") , 5)
    letters = [c for c in text if c.isalpha()]
    if len(letters) >= 3 and all(c.isupper() for c in letters):
        return 20  # an ALL-CAPS line
    return None


# A structural keyword alone on its line ("CAPITOLUL II", "Article 12"). Stricter than
# KEYWORD_PATTERNS, because it adds headings to a document that already has some.
STANDALONE_KEYWORD = re.compile(
    rf"^({'|'.join(HEADING_KEYWORDS)})\s*[:.]?\s*([IVXLCDM]+|\d+)$", re.I)
MIN_KEYWORD_HEADINGS = 3


ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}


def division_number(token: str) -> int:
    """ "12" -> 12, "IV" -> 4."""
    if token.isdigit():
        return int(token)
    values = [ROMAN[c] for c in token.upper()]
    return sum(-v if i + 1 < len(values) and v < values[i + 1] else v
               for i, v in enumerate(values))


def keyword_levels(lines: list[str]) -> dict[int, int]:
    """Line index -> heading level, for lines that are only a capitalised keyword and number.

    Divisions are numbered in order, so a line is accepted only if its number is higher
    than the last one at its rank; a new wider division (a chapter) restarts the
    narrower counters. This rejects keyword lines in tables and cross-references.
    """
    accepted: dict[int, int] = {}
    last: dict[int, int] = {}
    for i, line in enumerate(lines):
        text = line.strip()
        match = STANDALONE_KEYWORD.match(text) if text[:1].isupper() else None
        if not match:
            continue
        rank, number = pattern_rank(text), division_number(match.group(2))
        if number <= last.get(rank, 0):
            continue
        accepted[i] = rank
        last[rank] = number
        if rank != ANNEX_RANK:
            for narrower in [r for r in last if rank < r < ANNEX_RANK]:
                last[narrower] = 0
    # Annexes take the level of the widest division the document uses.
    ranks = sorted({rank for rank in accepted.values() if rank != ANNEX_RANK})
    order = {rank: level + 1 for level, rank in enumerate(ranks)}
    order[ANNEX_RANK] = 1
    return {i: order[rank] for i, rank in accepted.items()}


def pattern_levels(lines: list[str]) -> dict[int, int]:
    """Line index -> heading level, for lines whose wording marks them as headings."""
    ranked = {i: rank for i, line in enumerate(lines)
              if (rank := pattern_rank(line)) is not None}
    numbered = [i for i, rank in ranked.items() if 10 <= rank < 20]
    if len(numbered) < MIN_NUMBERED:
        ranked = {i: rank for i, rank in ranked.items() if not 10 <= rank < 20}
    order = {rank: min(level + 1, 6)
             for level, rank in enumerate(sorted(set(ranked.values())))}
    return {i: order[rank] for i, rank in ranked.items()}


def apply_pattern_headings(text: str) -> str:
    """Prefix "#" to the lines of a plain-text document that read like headings."""
    lines = text.split("\n")
    levels = pattern_levels(lines)
    return "\n".join(f"{'#' * levels[i]} {line.strip()}" if i in levels else line
                     for i, line in enumerate(lines))


# ----------------------------------------------------------------------- PDF

BOLD_FLAG = 1 << 4  # pymupdf span flag for bold
MARGIN_ZONE = 0.12  # top and bottom share of a page where running headers sit

# PyMuPDF also "finds" tables in boxed paragraphs (long cells) and in charts and forms
# (mostly empty cells). Kept as tables only: at most this share of empty cells, and
# cells this short on average. On the corpus this kept the scholarship, syllabus,
# ECTS grade and confusion-matrix tables and rejected every boxed paragraph and chart.
TABLE_MAX_EMPTY = 0.7
TABLE_MAX_CELL_CHARS = 40


def page_tables(page) -> list[tuple[pymupdf.Rect, float, str]]:
    """The page's tables as (area, top relative to the page height, text), top to bottom.

    Each row becomes one line "| a | b | c |". Extracted as plain text, a table loses its
    rows: every cell lands on its own line and a value no longer sits next to its label.
    Empty cells (merged or blank) are left out of the row.
    """
    height = page.rect.height or 1.0
    found = []
    for table in page.find_tables().tables:
        if table.row_count < 2 or table.col_count < 2:
            continue
        rows = [[" ".join((cell or "").split()) for cell in row] for row in table.extract()]
        cells = [cell for row in rows for cell in row]
        filled = [cell for cell in cells if cell]
        if (not filled or 1 - len(filled) / len(cells) > TABLE_MAX_EMPTY
                or sum(map(len, filled)) / len(filled) > TABLE_MAX_CELL_CHARS):
            continue
        area = pymupdf.Rect(table.bbox)
        found.append((area, area.y0 / height, "\n".join(table_lines(rows))))
    return sorted(found, key=lambda t: t[1])


NUMBER_CELL = re.compile(r"^[\d.,]+\s*%?$")


def table_lines(grid: list[list[str]]) -> list[str]:
    """A table as lines "| a | b |", with what a row needs to be read on its own.

    - Header: when the first row names at least three columns (and holds no numbers),
      every cell is written as "column: value", and a second header row under a
      merged cell ("Sursa de finanțare" over "Alocații bugetare" | "Venituri proprii") is
      joined to it. A small model otherwise could not tell which "x" or "-" is which.
    - Groups: in a table without such a header whose rows mostly end in a number (a
      scoring grid), a row with a
      single cell names a category ("Campionate europene"), and it is repeated in front of
      the rows below it. Otherwise "Locurile IV-VIII | 60" reads the same under every
      category, and the model took values from the wrong one.
    Empty cells (merged or blank) are left out.
    """
    rows = [row for row in grid if any(row)]
    header = None
    if (len(rows) >= 3 and sum(1 for c in rows[0] if c) >= 3
            and not any(NUMBER_CELL.match(c) for c in rows[0] if c)):
        header, rows = list(rows[0]), rows[1:]
        for i in range(1, len(header)):      # a merged cell covers the columns to its right
            header[i] = header[i] or header[i - 1]
        under = rows[0] if rows else []
        if under and not under[0] and not any(NUMBER_CELL.match(c) for c in under if c):
            header = [f"{h} – {s}" if s else h for h, s in zip(header, under)]
            rows = rows[1:]

    full = [[c for c in row if c] for row in rows]
    multi = [cells for cells in full if len(cells) >= 2]
    grouped = (header is None and bool(multi)
               and sum(bool(NUMBER_CELL.match(cells[-1])) for cells in multi) >= 0.6 * len(multi))

    lines = ["| " + " | ".join(dict.fromkeys(h for h in header if h)) + " |"] if header else []
    group = None
    for row, cells in zip(rows, full):
        if grouped and len(cells) == 1 and row[0]:
            # A category is a name; a cell ending in ":" or "." is a note or a blank to
            # fill in ("Alte activități: ......"), which ends the current category.
            group = None if cells[0].rstrip().endswith((":", ".")) else cells[0]
            lines.append(f"| {cells[0]} |")
            continue
        if header:
            cells = [f"{header[i]}: {c}" if header[i] and header[i] != c else c
                     for i, c in enumerate(row) if c]
        if grouped and group and len(cells) >= 2:
            cells = [group] + cells
        lines.append("| " + " | ".join(cells) + " |")
    return lines


def inside(bbox, areas) -> bool:
    """True if the centre of a line's box lies in one of the areas."""
    point = pymupdf.Point((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2)
    return any(point in area for area in areas)


def page_blocks(page, skip=()) -> list[list[tuple[str, float, bool, float]]]:
    """A page as blocks (roughly paragraphs) of lines:
    (text, largest font size, all bold, y position relative to the page height).
    Lines inside the `skip` areas (the page's tables) are left out."""
    height = page.rect.height or 1.0
    blocks = []
    for block in page.get_text("dict")["blocks"]:
        lines = []
        for line in block.get("lines", []):
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans or inside(line["bbox"], skip):
                continue
            lines.append(("".join(s["text"] for s in line["spans"]),
                          max(s["size"] for s in spans),
                          all(s["flags"] & BOLD_FLAG or "bold" in s["font"].lower()
                              for s in spans),
                          line["bbox"][1] / height))
        if lines:
            blocks.append(lines)
    return blocks


def in_margin(rel_y: float) -> bool:
    return rel_y < MARGIN_ZONE or rel_y > 1 - MARGIN_ZONE


def repeated_lines(pages) -> set[tuple[str, float]]:
    """Running headers and footers, as (text key, font size) pairs.

    A line counts if it sits in the top or bottom margin and repeats on most pages.
    The font size is part of the key so that a cover title is not removed along with
    the small running header that repeats it.
    """
    if len(pages) < 3:
        return set()
    counts: dict[tuple[str, float], int] = {}
    for blocks in pages:
        for key in {(repeat_key(t), round(size, 1))
                    for block in blocks for t, size, _, y in block if in_margin(y)}:
            counts[key] = counts.get(key, 0) + 1
    # At least two pages: a header usually starts on page 2, so in a short document
    # it may appear only twice.
    threshold = max(2, round(0.6 * len(pages)))
    return {key for key, n in counts.items() if n >= threshold}


def is_furniture(text: str, size: float, y: float, repeated: set[tuple[str, float]]) -> bool:
    """True for a running header or footer."""
    return in_margin(y) and (repeat_key(text), round(size, 1)) in repeated


def pdf_heading_level(pages, toc: dict[str, int], repeated: set[tuple[str, float]]):
    """A function (text, size, bold, y) -> heading level or None, for one PDF.

    Sources of headings, most trustworthy first: the PDF outline (bookmarks), if it
    matches the text; font size and weight; wording, for a PDF set in a single font.
    Wording is also added below the font headings when it finds far more structure
    than the fonts did (e.g. legal acts with chapters set in the body font).
    """
    lines = [line for blocks in pages for block in blocks for line in block]
    if toc:
        matched = sum(1 for text, _, _, _ in lines if norm(text) in toc)
        if matched >= len(toc) / 2:
            return lambda text, size, bold, y: (
                toc.get(norm(text)) if len(text.strip()) <= HEADING_MAX_CHARS else None)

    body = body_size([(text, size) for text, size, _, _ in lines])

    def is_heading(text: str, size: float, bold: bool, y: float) -> bool:
        return not is_furniture(text, size, y, repeated) \
            and looks_like_heading(text, size, bold, body)

    levels = level_map({round(size, 1) for text, size, bold, y in lines
                        if is_heading(text, size, bold, y)})

    if not levels:  # a single font throughout: use the wording
        content = [(text, y) for text, size, _, y in lines
                   if not is_furniture(text, size, y, repeated)]
        by_text = {norm(content[i][0]): level
                   for i, level in pattern_levels([t for t, _ in content]).items()}
        return lambda text, size, bold, y: by_text.get(norm(text))

    def level_of(text: str, size: float, bold: bool, y: float) -> int | None:
        return levels.get(round(size, 1)) if is_heading(text, size, bold, y) else None

    content = [text for text, size, _, y in lines if not is_furniture(text, size, y, repeated)]
    font_headings = sum(1 for text, size, bold, y in lines if is_heading(text, size, bold, y))
    by_keyword = keyword_levels(content)
    if len(by_keyword) >= MIN_KEYWORD_HEADINGS and len(by_keyword) > 2 * font_headings:
        depth = max(levels.values())
        # Keyed by (text, occurrence): the same text can be a heading once and a table
        # cell later. load_pdf asks about lines in document order, the order `content`
        # was built in, so counting occurrences lines the two up.
        seen: dict[str, int] = {}
        occurrence = []
        for text in content:
            key = norm(text)
            occurrence.append((key, seen.get(key, 0)))
            seen[key] = seen.get(key, 0) + 1
        below = {occurrence[i]: min(depth + level, 6) for i, level in by_keyword.items()}
        asked: dict[str, int] = {}

        def level_with_wording(text: str, size: float, bold: bool, y: float) -> int | None:
            key = norm(text)
            nth = asked.get(key, 0)
            asked[key] = nth + 1
            return level_of(text, size, bold, y) or below.get((key, nth))

        return level_with_wording

    return level_of


def load_pdf(path: Path, rel: str) -> Document:
    """One Document for the whole PDF, so sections spanning pages stay whole, with the
    character offset where each page starts, so a chunk can still cite its page."""
    with pymupdf.open(path) as pdf:
        tables = [page_tables(page) for page in pdf]
        pages = [page_blocks(page, [area for area, _, _ in on_page])
                 for page, on_page in zip(pdf, tables)]
        # Outline entries are (level, title, page); keyed by title, since the page a
        # bookmark points at can differ from where the heading is drawn.
        toc = {norm(title): min(level, 6)
               for level, title, _ in pdf.get_toc() if title.strip()}

    repeated = repeated_lines(pages)
    heading_level = pdf_heading_level(pages, toc, repeated)
    parts, page_starts, offset = [], [], 0
    for blocks, on_page in zip(pages, tables):
        rendered = []
        # A table goes in after the last block that starts above it. PyMuPDF does not
        # always list blocks top to bottom (a footnote can come first), so "before the
        # first block below it" could put a table ahead of its own title.
        after: dict[int, list[str]] = {}
        for _, top, table_text in on_page:
            above = [i for i, block in enumerate(blocks) if block[0][3] < top]
            after.setdefault(max(above) if above else -1, []).append(table_text)
        rendered.extend(after.get(-1, []))
        for index, block in enumerate(blocks):
            out = []
            for text, size, bold, y in block:
                if is_furniture(text, size, y, repeated):
                    continue
                level = heading_level(text, size, bold, y)
                out.append(f"{'#' * level} {text.strip()}" if level else text)
            if out:
                rendered.append("\n".join(out))
            rendered.extend(after.get(index, []))
        page_text = clean_text("\n\n".join(rendered), from_pdf=True)
        page_starts.append(offset)
        parts.append(page_text)
        offset += len(page_text) + 2  # the "\n\n" between pages
    text = "\n\n".join(parts)
    if len(text) < 50 * len(parts):
        print(f"[warn] {rel}: very little text, probably a scanned PDF (needs OCR)")
    return Document(sha(rel), rel, text,
                    {"format": "pdf", "num_pages": len(parts), "page_starts": page_starts})


# ---------------------------------------------------------------------- DOCX


def style_name(block) -> str:
    style = getattr(block, "style", None)
    return (style.name or "").lower() if style is not None else ""


def para_size(para, default: float) -> float:
    """Effective font size of a paragraph, in points."""
    sizes = [r.font.size.pt for r in para.runs if r.font.size]
    if sizes:
        return max(sizes)
    style = getattr(para, "style", None)
    if style is not None and style.font.size:
        return style.font.size.pt
    return default


def para_emphasized(para) -> bool:
    """True if every run with text is bold, or every one is underlined."""
    runs = [r for r in para.runs if r.text.strip()]
    return bool(runs) and (all(r.bold for r in runs) or all(r.underline for r in runs))


def docx_visual_levels(word, blocks) -> dict[int, int]:
    """Index in `blocks` -> heading level, for a Word file that uses no heading styles.

    Only called when no heading style is used at all; where styles are used, bold
    text is emphasis, not a title.
    """
    normal = word.styles["Normal"].font.size
    default = normal.pt if normal else 11.0

    paragraphs = []
    for i, block in enumerate(blocks):
        if isinstance(block, docx.table.Table) or style_name(block).startswith("list"):
            continue  # a short bold bullet is not a section title
        text = block.text.strip()
        if text:
            paragraphs.append((i, text, para_size(block, default), para_emphasized(block)))

    body = body_size([(text, size) for _, text, size, _ in paragraphs])
    candidates = {i: size for i, text, size, emph in paragraphs
                  if looks_like_heading(text, size, emph, body)}
    if candidates:
        levels = level_map({round(size, 1) for size in candidates.values()})
        return {i: levels[round(size, 1)] for i, size in candidates.items()}

    # No formatting to go on either: use the wording.
    by_line = pattern_levels([text for _, text, _, _ in paragraphs])
    return {paragraphs[i][0]: level for i, level in by_line.items()}


def load_docx(path: Path, rel: str) -> Document:
    """Word to markdown: headings -> "#", lists -> "- ", tables -> "| a | b |",
    with paragraphs and tables kept in their original order."""
    word = docx.Document(str(path))
    blocks = list(word.iter_inner_content())
    marked = [b for b in blocks if style_name(b).startswith(("title", "heading"))]
    # With a "Title" paragraph, Heading 1 sits one level below it.
    shift = 1 if any(style_name(b) == "title" for b in blocks) else 0
    visual = {} if marked else docx_visual_levels(word, blocks)
    lines = []
    for index, block in enumerate(blocks):
        if isinstance(block, docx.table.Table):
            for i, row in enumerate(block.rows):
                values = [c.text.strip().replace("\n", " ") for c in row.cells]
                # A merged cell is returned once per column it spans; keep one copy.
                cells = [v if j == 0 or v != values[j - 1] else ""
                         for j, v in enumerate(values)]
                lines.append("| " + " | ".join(cells) + " |")
                if i == 0:  # markdown header separator
                    lines.append("|" + " --- |" * len(cells))
            lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = style_name(block)
        if style == "title":
            lines.append(f"# {text}")
        elif style.startswith("heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lines.append(f"{'#' * min(level + shift, 6)} {text}")
        elif style.startswith("list"):
            lines.append(f"- {text}")
        elif index in visual:  # a title typed in bold instead of styled
            lines.append(f"{'#' * visual[index]} {text}")
        else:
            lines.append(text)
        lines.append("")
    title = word.core_properties.title or None
    return Document(sha(rel), rel, clean_text("\n".join(lines)), {"format": "docx", "title": title})


# ------------------------------------------------------------- other formats


def rows_to_markdown(rows: list[list]) -> str:
    """Rows as a markdown table, without columns that are empty in every row and without
    the empty tail of each row.

    Spreadsheets laid out as forms are mostly empty cells, which would otherwise fill
    chunks with pipes and spaces. Empty cells between values are kept, so values stay
    under their column header.
    """
    rows = [["" if c is None else str(c).strip().replace("\n", " ") for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    used = [j for j in range(width) if any(j < len(r) and r[j] for r in rows)]
    out = []
    for i, row in enumerate(rows):
        cells = [row[j] if j < len(row) else "" for j in used]
        while len(cells) > 1 and not cells[-1]:
            cells.pop()
        out.append("| " + " | ".join(cells) + " |")
        if i == 0:
            out.append("|" + " --- |" * len(cells))
    return "\n".join(out)


def load_csv(path: Path, rel: str) -> Document:
    raw = read_text_file(path, rel)
    try:
        delimiter = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    return Document(sha(rel), rel, clean_text(rows_to_markdown(rows)),
                    {"format": "csv", "num_rows": len(rows)})


def load_pptx(path: Path, rel: str) -> Document:
    """One "#" section per slide, with its bullets and speaker notes."""
    deck = pptx.Presentation(str(path))
    lines = []
    for number, slide in enumerate(deck.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        lines += [f"# {title or f'Slide {number}'}", ""]
        for shape in slide.shapes:
            # Compared by id: python-pptx returns a new proxy object on every access, so
            # `shape is slide.shapes.title` is never true and the title was repeated.
            if not shape.has_text_frame or (title_shape is not None
                                            and shape.shape_id == title_shape.shape_id):
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    lines.append(f"{'  ' * (para.level or 0)}- {text}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines += ["", f"Speaker notes: {notes}"]
        lines.append("")
    return Document(sha(rel), rel, clean_text("\n".join(lines)),
                    {"format": "pptx", "num_slides": len(deck.slides)})


def load_xlsx(path: Path, rel: str) -> Document:
    """One "#" section per sheet, holding its table. Formulas are read as their values."""
    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        lines = []
        for sheet in book.worksheets:
            table = rows_to_markdown([list(row) for row in sheet.iter_rows(values_only=True)])
            if table:
                lines += [f"# {sheet.title}", "", table, ""]
        sheets = len(book.worksheets)
    finally:
        book.close()
    return Document(sha(rel), rel, clean_text("\n".join(lines)),
                    {"format": "xlsx", "num_sheets": sheets})


# ------------------------------------------------------------------- entry points


def page_at(page_starts: list[int], char_pos: int) -> int:
    """1-based page number of a character position in a PDF Document."""
    return bisect_right(page_starts, char_pos)


LOADERS = {".md": load_text, ".markdown": load_text, ".txt": load_text,
           ".html": load_html, ".pdf": load_pdf, ".docx": load_docx,
           ".pptx": load_pptx, ".xlsx": load_xlsx, ".csv": load_csv}


def load_file(path: Path, rel: str) -> Document:
    """Read one file with the loader for its extension.

    `rel` is the document's name ("folder/file.pdf"); `path` is where its bytes are
    now (an upload is read from a temporary file). Raises ValueError for an
    unsupported format, or when no text could be read (e.g. a scanned PDF).
    """
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"unsupported format {path.suffix or '(no extension)'}; "
                         f"supported: {', '.join(sorted(LOADERS))}")
    doc = loader(path, rel)
    if not doc.text.strip():
        raise ValueError("no text could be read from this file "
                         "(a scanned PDF needs OCR, which is not built in)")
    doc.metadata["content_hash"] = sha(doc.text)
    doc.metadata["num_chars"] = len(doc.text)
    return doc


def load_directory(raw_dir: Path) -> list[Document]:
    """Every supported file under `raw_dir`; unreadable files are reported and skipped."""
    docs = []
    for path in sorted(raw_dir.rglob("*")):
        if path.suffix.lower() not in LOADERS or not path.is_file():
            continue
        rel = path.relative_to(raw_dir).as_posix()
        try:
            docs.append(load_file(path, rel))
        except Exception as e:
            print(f"[skip] {rel}: {e}")
    return docs


def save_jsonl(docs: list[Document], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        for d in docs:
            f.write(json.dumps(asdict(d), ensure_ascii=False) + "\n")
